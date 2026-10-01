from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUTPUT_DIR = Path("generated_documents")
OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


# ============================================================
# TEXT FILE
# ============================================================

def save_text_file(filename: str, content: str) -> str:

    path = OUTPUT_DIR / filename

    path.write_text(
        content,
        encoding="utf-8"
    )

    return str(path)


# ============================================================
# DOCX
# ============================================================

def create_docx(
    filename: str,
    title: str,
    content: str
) -> str:

    path = OUTPUT_DIR / filename

    document = Document()

    # LegalEase header
    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = header.add_run("⚖ LegalEase")
    run.bold = True
    run.font.size = Pt(18)

    # Title
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph()

    # Content
    for line in content.split("\n"):

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)
        run.font.size = Pt(10.5)

        if (
            line.startswith("WITNESSETH")
            or line.startswith("WHEREAS")
            or line.startswith("NOW, THEREFORE")
            or line.startswith("1. ")
            or line.startswith("2. ")
            or line.startswith("3. ")
            or line.startswith("4. ")
            or line.startswith("5. ")
            or line.startswith("6. ")
            or line.startswith("7. ")
        ):
            run.bold = True

    # Footer
    section = document.sections[0]

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.add_run(
        "LegalEase Inc. | contact@legalease.com | All Rights Reserved"
    )

    footer_run.font.size = Pt(8)

    document.save(path)

    return str(path)


# ============================================================
# PDF
# ============================================================

def create_pdf(
    filename: str,
    title: str,
    content: str
) -> str:

    path = OUTPUT_DIR / filename

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=25 * mm,
        leftMargin=25 * mm,
        topMargin=20 * mm,
        bottomMargin=22 * mm
    )

    styles = getSampleStyleSheet()

    logo_style = ParagraphStyle(
        "Logo",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        spaceAfter=4
    )

    title_style = ParagraphStyle(
        "Title",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=20,
        spaceAfter=18
    )

    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        spaceAfter=6
    )

    heading_style = ParagraphStyle(
        "Heading",
        parent=body_style,
        fontName="Helvetica-Bold",
        spaceBefore=4,
        spaceAfter=6
    )

    story = []

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "⚖ LegalEase",
            logo_style
        )
    )

    story.append(
        Paragraph(
            title,
            title_style
        )
    )

    # --------------------------------------------------------
    # CONTENT
    # --------------------------------------------------------

    for line in content.split("\n"):

        line = line.strip()

        if not line:

            story.append(
                Spacer(1, 4)
            )

            continue

        # Escape HTML characters
        safe_line = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        # Bold legal headings
        if (
            safe_line.startswith("WITNESSETH")
            or safe_line.startswith("WHEREAS")
            or safe_line.startswith("NOW, THEREFORE")
            or safe_line.startswith("1. ")
            or safe_line.startswith("2. ")
            or safe_line.startswith("3. ")
            or safe_line.startswith("4. ")
            or safe_line.startswith("5. ")
            or safe_line.startswith("6. ")
            or safe_line.startswith("7. ")
        ):

            story.append(
                Paragraph(
                    f"<b>{safe_line}</b>",
                    heading_style
                )
            )

        elif safe_line.startswith("- "):

            story.append(
                Paragraph(
                    "• " + safe_line[2:],
                    body_style
                )
            )

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style
                )
            )

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    def add_footer(canvas, doc):

        canvas.saveState()

        width, height = A4

        canvas.setFont(
            "Helvetica",
            8
        )

        canvas.drawCentredString(
            width / 2,
            12 * mm,
            "LegalEase Inc. | contact@legalease.com | All Rights Reserved"
        )

        canvas.restoreState()

    document.build(
        story,
        onFirstPage=add_footer,
        onLaterPages=add_footer
    )

    return str(path)