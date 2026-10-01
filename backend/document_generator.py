from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)
from reportlab.pdfbase.pdfmetrics import stringWidth

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# OUTPUT FOLDER
# ============================================================

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "generated_documents"

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


# ============================================================
# TEXT FILE
# ============================================================

def save_text_file(
    filename: str,
    content: str
) -> str:

    path = OUTPUT_DIR / filename

    path.write_text(
        content,
        encoding="utf-8"
    )

    return str(path)


# ============================================================
# DOCX
# ============================================================

def create_docx(
    filename: str,
    title: str,
    content: str
) -> str:

    path = OUTPUT_DIR / filename

    document = Document()

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    header = document.add_paragraph()

    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = header.add_run(
        "⚖ LegalEase"
    )

    run.bold = True
    run.font.size = Pt(20)

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    heading = document.add_paragraph()

    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = heading.add_run(
        title
    )

    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph("")

    # --------------------------------------------------------
    # CONTENT
    # --------------------------------------------------------

    for line in content.split("\n"):

        line = line.strip()

        if not line:

            document.add_paragraph()

            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            line
        )

        run.font.size = Pt(10.5)

        # Bold legal headings

        if (
            line.startswith("WITNESSETH")
            or line.startswith("WHEREAS")
            or line.startswith("NOW, THEREFORE")
            or line.startswith("Between:")
            or line.startswith("And:")
            or line.startswith("1. ")
            or line.startswith("2. ")
            or line.startswith("3. ")
            or line.startswith("4. ")
            or line.startswith("5. ")
            or line.startswith("6. ")
            or line.startswith("7. ")
        ):

            run.bold = True

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    section = document.sections[0]

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "LegalEase Inc. | "
        "AI-Powered Legal Document Assistant"
    )

    footer_run.font.size = Pt(8)

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    document.save(path)

    return str(path)


# ============================================================
# PDF LOGO
# ============================================================

def draw_legal_ease_logo(
    pdf,
    center_x,
    top_y
):

    pdf.setStrokeColor(
        colors.black
    )

    pdf.setLineWidth(2)

    # --------------------------------------------------------
    # CENTER POLE
    # --------------------------------------------------------

    pdf.line(
        center_x,
        top_y,
        center_x,
        top_y - 45
    )

    # --------------------------------------------------------
    # HORIZONTAL BEAM
    # --------------------------------------------------------

    pdf.line(
        center_x - 48,
        top_y - 8,
        center_x + 48,
        top_y - 8
    )

    # --------------------------------------------------------
    # LEFT SUPPORT
    # --------------------------------------------------------

    pdf.line(
        center_x - 48,
        top_y - 8,
        center_x - 60,
        top_y - 30
    )

    pdf.line(
        center_x - 48,
        top_y - 8,
        center_x - 36,
        top_y - 30
    )

    # --------------------------------------------------------
    # LEFT PAN
    # --------------------------------------------------------

    pdf.line(
        center_x - 60,
        top_y - 30,
        center_x - 36,
        top_y - 30
    )

    # --------------------------------------------------------
    # RIGHT SUPPORT
    # --------------------------------------------------------

    pdf.line(
        center_x + 48,
        top_y - 8,
        center_x + 36,
        top_y - 30
    )

    pdf.line(
        center_x + 48,
        top_y - 8,
        center_x + 60,
        top_y - 30
    )

    # --------------------------------------------------------
    # RIGHT PAN
    # --------------------------------------------------------

    pdf.line(
        center_x + 36,
        top_y - 30,
        center_x + 60,
        top_y - 30
    )

    # --------------------------------------------------------
    # BOTTOM STAND
    # --------------------------------------------------------

    pdf.line(
        center_x,
        top_y - 45,
        center_x,
        top_y - 55
    )

    pdf.line(
        center_x - 20,
        top_y - 55,
        center_x + 20,
        top_y - 55
    )


# ============================================================
# PDF
# ============================================================

def create_pdf(
    filename: str,
    title: str,
    content: str
) -> str:

    path = OUTPUT_DIR / filename

    # --------------------------------------------------------
    # DOCUMENT
    # --------------------------------------------------------

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=22 * mm,
        leftMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=22 * mm
    )

    styles = getSampleStyleSheet()

    # --------------------------------------------------------
    # LOGO STYLE
    # --------------------------------------------------------

    logo_style = ParagraphStyle(
        "LegalEaseLogo",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.black,
        spaceAfter=5
    )

    # --------------------------------------------------------
    # TITLE STYLE
    # --------------------------------------------------------

    title_style = ParagraphStyle(
        "DocumentTitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=21,
        textColor=colors.black,
        spaceAfter=16
    )

    # --------------------------------------------------------
    # BODY STYLE
    # --------------------------------------------------------

    body_style = ParagraphStyle(
        "DocumentBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        textColor=colors.black,
        spaceAfter=6
    )

    # --------------------------------------------------------
    # HEADING STYLE
    # --------------------------------------------------------

    heading_style = ParagraphStyle(
        "DocumentHeading",
        parent=body_style,
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=15,
        spaceBefore=5,
        spaceAfter=6
    )

    # --------------------------------------------------------
    # STORY
    # --------------------------------------------------------

    story = []

    # --------------------------------------------------------
    # LEGALEASE HEADER
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "⚖ LegalEase",
            logo_style
        )
    )

    story.append(
        Paragraph(
            title,
            title_style
        )
    )

    # --------------------------------------------------------
    # HEADER LINE
    # --------------------------------------------------------

    story.append(
        Paragraph(
            "<hr/>",
            body_style
        )
    )

    story.append(
        Spacer(
            1,
            8
        )
    )

    # --------------------------------------------------------
    # CONTENT
    # --------------------------------------------------------

    for line in content.split("\n"):

        line = line.strip()

        # Empty line

        if not line:

            story.append(
                Spacer(
                    1,
                    5
                )
            )

            continue

        # ----------------------------------------------------
        # ESCAPE HTML
        # ----------------------------------------------------

        safe_line = (
            line
            .replace(
                "&",
                "&amp;"
            )
            .replace(
                "<",
                "&lt;"
            )
            .replace(
                ">",
                "&gt;"
            )
        )

        # ----------------------------------------------------
        # HEADINGS
        # ----------------------------------------------------

        is_heading = (
            safe_line.startswith("WITNESSETH")
            or safe_line.startswith("WHEREAS")
            or safe_line.startswith("NOW, THEREFORE")
            or safe_line.startswith("Between:")
            or safe_line.startswith("And:")
            or safe_line.startswith("1. ")
            or safe_line.startswith("2. ")
            or safe_line.startswith("3. ")
            or safe_line.startswith("4. ")
            or safe_line.startswith("5. ")
            or safe_line.startswith("6. ")
            or safe_line.startswith("7. ")
        )

        if is_heading:

            story.append(
                Paragraph(
                    f"<b>{safe_line}</b>",
                    heading_style
                )
            )

        elif safe_line.startswith("- "):

            story.append(
                Paragraph(
                    "• "
                    + safe_line[2:],
                    body_style
                )
            )

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style
                )
            )

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    def add_footer(
        canvas,
        doc
    ):

        canvas.saveState()

        width, height = A4

        canvas.setFont(
            "Helvetica",
            8
        )

        canvas.setFillColor(
            colors.grey
        )

        canvas.drawCentredString(
            width / 2,
            10 * mm,
            "LegalEase Inc. | "
            "AI-Powered Legal Document Assistant | "
            "For educational and informational purposes only."
        )

        canvas.restoreState()

    # --------------------------------------------------------
    # BUILD PDF
    # --------------------------------------------------------

    document.build(
        story,
        onFirstPage=add_footer,
        onLaterPages=add_footer
    )

    return str(path)